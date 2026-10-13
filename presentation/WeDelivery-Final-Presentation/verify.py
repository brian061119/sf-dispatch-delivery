# -*- coding: utf-8 -*-
"""Layout guard: estimate rendered text height and flag overflow / out-of-canvas shapes."""
import math
import os
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, 'WeDelivery-Final-Presentation.pptx')
EMU_PER_PX = 9525.0
FONTS = {
    'segoe ui': 'C:/Windows/Fonts/segoeui.ttf',
    'segoe ui semibold': 'C:/Windows/Fonts/seguisb.ttf',
    'segoe ui bold': 'C:/Windows/Fonts/segoeuib.ttf',
    'consolas': 'C:/Windows/Fonts/consola.ttf',
}
_cache = {}


def font(name, pt, bold):
    key = (name, pt, bold)
    if key in _cache:
        return _cache[key]
    path = FONTS.get((name or 'segoe ui').lower())
    if bold and (name or '').lower() == 'segoe ui':
        path = FONTS['segoe ui bold']
    px = int(round(pt * 96.0 / 72.0))
    try:
        f = ImageFont.truetype(path, px)
    except Exception:
        f = ImageFont.load_default()
    _cache[key] = f
    return f


def text_width(s, name, pt, bold):
    f = font(name, pt, bold)
    return f.getlength(s)


def est_paragraph_height(p, inner_w):
    """Return estimated height in px for one paragraph."""
    sizes, texts = [], []
    for r in p.runs:
        txt = r.text or ''
        pt = r.font.size.pt if r.font.size else 18.0
        nm = r.font.name or 'Segoe UI'
        bd = bool(r.font.bold)
        texts.append((txt, nm, pt, bd))
    if not texts:
        return 0.0
    total = 0.0
    line_h_max = 0.0
    # rough wrap model: accumulate width, break at inner_w
    cur = 0.0
    lines = 1
    for txt, nm, pt, bd in texts:
        # PowerPoint "single" line height is ~1.2x the font size
        line_h_max = max(line_h_max, pt * 96.0 / 72.0 * 1.2)
        words = txt.split(' ')
        for i, w in enumerate(words):
            piece = w if i == 0 else ' ' + w
            ww = text_width(piece, nm, pt, bd)
            if cur + ww > inner_w and cur > 0:
                lines += 1
                cur = ww
            else:
                cur += ww
    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
    total = lines * line_h_max * ls
    sa = p.space_after.pt if p.space_after else 0.0
    sb = p.space_before.pt if p.space_before else 0.0
    return total + (sa + sb) * 96.0 / 72.0


def main():
    prs = Presentation(PPTX)
    SW = prs.slide_width / EMU_PER_PX
    SH = prs.slide_height / EMU_PER_PX
    problems = []

    for idx, slide in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            left = shp.left / EMU_PER_PX if shp.left is not None else 0
            top = shp.top / EMU_PER_PX if shp.top is not None else 0
            w = shp.width / EMU_PER_PX if shp.width is not None else 0
            h = shp.height / EMU_PER_PX if shp.height is not None else 0

            if left < -1 or top < -1 or left + w > SW + 1 or top + h > SH + 1:
                problems.append(
                    'S%02d OUT-OF-CANVAS %s  box=(%.0f,%.0f %.0fx%.0f)'
                    % (idx, shp.shape_type, left, top, w, h))

            if not shp.has_text_frame:
                continue
            tf = shp.text_frame
            if not (tf.text or '').strip():
                continue
            ml = (tf.margin_left or 0) / EMU_PER_PX
            mr = (tf.margin_right or 0) / EMU_PER_PX
            mt = (tf.margin_top or 0) / EMU_PER_PX
            mb = (tf.margin_bottom or 0) / EMU_PER_PX
            inner_w = max(10.0, w - ml - mr)
            inner_h = h - mt - mb
            need = sum(est_paragraph_height(p, inner_w) for p in tf.paragraphs)
            if need > inner_h + 6:
                sample = (tf.text or '').replace('\n', ' / ')[:70]
                problems.append(
                    'S%02d OVERFLOW  need~%.0fpx have %.0fpx  @(%.0f,%.0f %.0fx%.0f)  "%s"'
                    % (idx, need, inner_h, left, top, w, h, sample))

    # tables: verify total column width and estimated row heights
    for idx, slide in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            if not shp.has_table:
                continue
            tbl = shp.table
            colsum = sum(c.width for c in tbl.columns) / EMU_PER_PX
            tblh = shp.height / EMU_PER_PX
            rowh = sum(r.height for r in tbl.rows) / EMU_PER_PX
            if abs(colsum - shp.width / EMU_PER_PX) > 2:
                problems.append('S%02d TABLE cols=%.0f vs width=%.0f'
                                % (idx, colsum, shp.width / EMU_PER_PX))
            if rowh > tblh + 2:
                problems.append('S%02d TABLE rows total=%.0f > box %.0f'
                                % (idx, rowh, tblh))
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row.cells):
                    cw = tbl.columns[ci].width / EMU_PER_PX - 24
                    ch = row.height / EMU_PER_PX
                    need = sum(est_paragraph_height(p, max(10.0, cw))
                               for p in cell.text_frame.paragraphs)
                    if need > ch + 4:
                        problems.append(
                            'S%02d TABLE r%dc%d need~%.0f > row %.0f  "%s"'
                            % (idx, ri, ci, need, ch, cell.text[:48]))

    print('slides:', len(prs.slides._sldIdLst))
    if problems:
        print('PROBLEMS (%d):' % len(problems))
        for p in problems:
            print('  -', p)
    else:
        print('NO PROBLEMS')


main()
